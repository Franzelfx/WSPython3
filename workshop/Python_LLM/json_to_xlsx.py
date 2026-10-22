"""
Wandelt die EV-Produktionsdaten (JSON) in eine XLSX-Datei um.

Benötigt: pip install openpyxl
Aufruf:   python json_to_xlsx.py
          python json_to_xlsx.py eingabe.json ausgabe.xlsx
"""

import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASIS = Path(__file__).parent
EINGABE = Path(sys.argv[1]) if len(sys.argv) > 1 else BASIS / "ev_produktion.json"
AUSGABE = Path(sys.argv[2]) if len(sys.argv) > 2 else BASIS / "ev_produktion.xlsx"

KOPF_FARBE = PatternFill("solid", fgColor="1F4E78")
KOPF_SCHRIFT = Font(bold=True, color="FFFFFF")


def lade_json(pfad: Path) -> dict:
    with pfad.open(encoding="utf-8") as datei:
        return json.load(datei)


def schreibe_kopfzeile(blatt, spalten: list[str]) -> None:
    blatt.append(spalten)
    for zelle in blatt[1]:
        zelle.fill = KOPF_FARBE
        zelle.font = KOPF_SCHRIFT
        zelle.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    blatt.freeze_panes = "A2"


def passe_spaltenbreite_an(blatt) -> None:
    for spalte in blatt.columns:
        laenge = max(len(str(zelle.value)) if zelle.value is not None else 0 for zelle in spalte)
        blatt.column_dimensions[get_column_letter(spalte[0].column)].width = min(laenge + 3, 40)


def erstelle_rohdaten(wb: Workbook, eintraege: list[dict]) -> None:
    blatt = wb.active
    blatt.title = "Rohdaten"
    spalten = ["Jahr", "Zulieferer", "Werk", "Land", "Region", "Produkt",
               "Produziert (Stück)", "Kapazität (Stück)", "Auslastung"]
    schreibe_kopfzeile(blatt, spalten)

    for zeile, e in enumerate(eintraege, start=2):
        blatt.append([
            e["jahr"], e["zulieferer"], e["werk"], e["land"], e["region"], e["produkt"],
            e["produziert"], e["kapazitaet"],
            # Auslastung als Excel-Formel, damit sie sich bei Datenänderungen mitaktualisiert
            f"=G{zeile}/H{zeile}",
        ])

    letzte_zeile = blatt.max_row
    for zeile in range(2, letzte_zeile + 1):
        blatt[f"G{zeile}"].number_format = "#,##0"
        blatt[f"H{zeile}"].number_format = "#,##0"
        blatt[f"I{zeile}"].number_format = "0.0%"

    blatt.auto_filter.ref = f"A1:I{letzte_zeile}"
    passe_spaltenbreite_an(blatt)


def erstelle_uebersicht(wb: Workbook, eintraege: list[dict]) -> None:
    """Summe der produzierten Stückzahlen pro Jahr und Region."""
    summen = defaultdict(lambda: defaultdict(int))
    jahre = sorted({e["jahr"] for e in eintraege})
    regionen = sorted({e["region"] for e in eintraege})

    for e in eintraege:
        summen[e["region"]][e["jahr"]] += e["produziert"]

    blatt = wb.create_sheet("Übersicht")
    schreibe_kopfzeile(blatt, ["Region"] + [str(j) for j in jahre] + ["Gesamt"])

    for region in regionen:
        werte = [summen[region][j] for j in jahre]
        blatt.append([region] + werte + [sum(werte)])

    # Gesamtzeile
    gesamt = [sum(summen[r][j] for r in regionen) for j in jahre]
    blatt.append(["Gesamt"] + gesamt + [sum(gesamt)])
    for zelle in blatt[blatt.max_row]:
        zelle.font = Font(bold=True)

    for zeile in blatt.iter_rows(min_row=2, min_col=2):
        for zelle in zeile:
            zelle.number_format = "#,##0"
    passe_spaltenbreite_an(blatt)


def main() -> None:
    daten = lade_json(EINGABE)
    eintraege = daten["eintraege"]

    wb = Workbook()
    erstelle_rohdaten(wb, eintraege)
    erstelle_uebersicht(wb, eintraege)
    wb.save(AUSGABE)

    print(f"{len(eintraege)} Datensätze gespeichert unter: {AUSGABE}")


if __name__ == "__main__":
    main()
